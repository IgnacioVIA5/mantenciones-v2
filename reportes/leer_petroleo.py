#!/usr/bin/env python3
"""
Lee el Excel CONTROL PETROLEO y calcula, por patente, la lectura mas reciente
de horometro y odometro (de forma independiente), comparandola contra el
estado actual en Firestore (equipos.json, generado por fetch_flota.mjs).

Uso:
    python3 leer_petroleo.py "/ruta/a/CONTROL PETROLEO 2026.xlsx"

Genera lecturas_pendientes.json con el plan de actualizacion para que
aplicar_lecturas.mjs lo suba a Firestore.
"""
import sys
import re
import json
import datetime
import openpyxl
from pathlib import Path

HOJAS = ["ESTANQUE 1- CASA", "ESTANQUE 2 - PLANTA"]
COL_FECHA, COL_MAQUINARIA, COL_PATENTE, COL_HOROMETRO, COL_ODOMETRO = 1, 4, 5, 8, 9

# Multiplo maximo de crecimiento aceptado sin revision manual (filtra
# errores de columna, ej. odometro tipeado en la celda de horometro).
FACTOR_SOSPECHA = 5


def norm(p):
    if not p:
        return None
    return re.sub(r"[\s\-]", "", str(p)).upper()


def leer_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    filas = []
    for hoja in HOJAS:
        ws = wb[hoja]
        for r in range(3, ws.max_row + 1):
            fecha = ws.cell(row=r, column=COL_FECHA).value
            patente = ws.cell(row=r, column=COL_PATENTE).value
            if not patente or not isinstance(fecha, datetime.datetime):
                continue
            filas.append({
                "fecha": fecha.date().isoformat(),
                "patente_norm": norm(patente),
                "horometro": ws.cell(row=r, column=COL_HOROMETRO).value,
                "odometro": ws.cell(row=r, column=COL_ODOMETRO).value,
                "hoja": hoja,
            })
    return filas


def ultimo_valor(filas, patente_norm, campo):
    """Ultima fila (por fecha) con valor no nulo en `campo`, para esa patente."""
    candidatas = [f for f in filas if f["patente_norm"] == patente_norm and f[campo] is not None]
    if not candidatas:
        return None
    candidatas.sort(key=lambda f: f["fecha"])
    ultima = candidatas[-1]
    return {"valor": ultima[campo], "fecha": ultima["fecha"]}


def plan_para_equipo(e, filas):
    patente = e.get("patente")
    pn = norm(patente)
    categoria = e.get("categoria")
    items = []

    if categoria == "CAMION":
        objetivos = [
            ("horometro", "horaActual", "horaActualFecha", e.get("horaActual"), e.get("horaActualFecha")),
            ("odometro", "odometro", "odometroFecha", e.get("odometro"), e.get("odometroFecha")),
        ]
    elif categoria == "CAMIONETA":
        # La camioneta solo registra un valor (km) en horaActual; en el excel
        # puede venir en la columna Odometro u, ocasionalmente por error de
        # tipeo, en Horometro. Se usa Odometro y si falta, Horometro.
        candidatas = [f for f in filas if f["patente_norm"] == pn and (f["odometro"] is not None or f["horometro"] is not None)]
        if not candidatas:
            return items
        candidatas.sort(key=lambda f: f["fecha"])
        ultima = candidatas[-1]
        valor = ultima["odometro"] if ultima["odometro"] is not None else ultima["horometro"]
        objetivos = [(None, "horaActual", "horaActualFecha", e.get("horaActual"), e.get("horaActualFecha"))]
        items.append(_evaluar(patente, categoria, "horaActual (km)", "horaActual", "horaActualFecha",
                               e.get("horaActual"), e.get("horaActualFecha"), valor, ultima["fecha"]))
        return items
    elif categoria in ("CARGADOR", "EXCAVADORA", "GENERADOR"):
        objetivos = [
            ("horometro", "horaActual", "horaActualFecha", e.get("horaActual"), e.get("horaActualFecha")),
        ]
    else:
        return items

    for campo_excel, campo_fs, campo_fecha_fs, actual, actual_fecha in objetivos:
        nuevo = ultimo_valor(filas, pn, campo_excel)
        if nuevo is None:
            continue
        label = "horaActual (hr)" if campo_fs == "horaActual" else "odometro (km)"
        items.append(_evaluar(patente, categoria, label, campo_fs, campo_fecha_fs,
                               actual, actual_fecha, nuevo["valor"], nuevo["fecha"]))
    return items


def _evaluar(patente, categoria, label, campo_fs, campo_fecha_fs, actual, actual_fecha, valor_nuevo, fecha_nueva):
    actual_num = actual if isinstance(actual, (int, float)) else 0
    valor_nuevo = round(float(valor_nuevo))

    if actual_num > 0 and valor_nuevo > actual_num * FACTOR_SOSPECHA:
        accion = "revisar"
        motivo = f"salto sospechoso: {valor_nuevo} es >{FACTOR_SOSPECHA}x el valor actual ({actual_num}); posible error de columna en el excel"
    elif valor_nuevo > actual_num:
        accion = "actualizar"
        motivo = ""
    else:
        accion = "sin_cambios"
        motivo = "el excel no trae una lectura mas nueva que la ya registrada"

    return {
        "patente": patente,
        "categoria": categoria,
        "campo": label,
        "campoFirestore": campo_fs,
        "campoFechaFirestore": campo_fecha_fs,
        "valorActual": actual_num,
        "fechaActual": actual_fecha,
        "valorNuevo": valor_nuevo,
        "fechaNueva": fecha_nueva,
        "accion": accion,
        "motivo": motivo,
    }


def main():
    if len(sys.argv) < 2:
        print("Uso: python3 leer_petroleo.py <ruta al xlsx>")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    here = Path(__file__).parent
    equipos = json.loads((here / "equipos.json").read_text())

    filas = leer_excel(xlsx_path)
    patentes_excel = {f["patente_norm"] for f in filas}
    patentes_app = {norm(e["patente"]): e["patente"] for e in equipos if e.get("patente")}

    plan = []
    for e in equipos:
        if not e.get("patente"):
            continue
        plan.extend(plan_para_equipo(e, filas))

    no_en_app = sorted(patentes_excel - set(patentes_app.keys()))
    no_en_excel = sorted(set(patentes_app.keys()) - patentes_excel)

    salida = {
        "generadoEl": datetime.date.today().isoformat(),
        "archivoOrigen": str(xlsx_path),
        "plan": plan,
        "patentesExcelSinApp": no_en_app,
        "patentesAppSinExcel": [patentes_app[p] for p in no_en_excel],
    }

    out_path = here / "lecturas_pendientes.json"
    out_path.write_text(json.dumps(salida, ensure_ascii=False, indent=2))

    print(f"\n== Plan de actualizacion ({out_path.name}) ==\n")
    for item in plan:
        if item["accion"] == "sin_cambios":
            continue
        marca = "⚠️ REVISAR" if item["accion"] == "revisar" else "✅ actualizar"
        print(f"{item['patente']:10} {item['campo']:16} {item['valorActual']} ({item['fechaActual']}) -> "
              f"{item['valorNuevo']} ({item['fechaNueva']})  [{marca}]"
              + (f"  -- {item['motivo']}" if item['motivo'] else ""))

    if no_en_app:
        print(f"\nPatentes en el excel que no estan en la app (se ignoran): {no_en_app}")
    if no_en_excel:
        print(f"Patentes en la app sin filas en el excel (sin cambios): {no_en_excel}")


if __name__ == "__main__":
    main()
